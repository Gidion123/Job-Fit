"""Import Dion's D-080 gap relevance labels into a new versioned gold bundle (r4).

Default is a dry run that only validates and prints the counts. `--write` creates
`evals/gold/development_v13_reviewed_20261004_gap_r4/`, run only after Dion
approves the import in chat. Bundle r3 is copied byte for byte and never edited;
only `relevance_gold.jsonl` gains the new rows, and a new manifest records hashes.

Checks: workbook closed (no Excel lock file), every manifest item present once
with the same CV and job, relevance in 0..3, a non-empty reason, development
split only, no pair already in r3, no extra rows.
"""
from __future__ import annotations

import argparse
from collections import Counter
from hashlib import sha256
import json
from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
GAP = ROOT / 'evals/labeling/dev_relevance_gap_20261004_v2'
WORKBOOK = GAP / 'JobFit_Dev_Relevance_Gap_v2.xlsx'
BASE = ROOT / 'evals/gold/development_v13_reviewed_20261003_stage1_r3'
TARGET = ROOT / 'evals/gold/development_v13_reviewed_20261004_gap_r4'


def digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def read_rows() -> list[dict]:
    from openpyxl import load_workbook
    if any(GAP.glob('~$*')):
        raise SystemExit('Close the workbook in Excel first (lock file present)')
    ws = load_workbook(WORKBOOK, read_only=True, data_only=True)['Items']
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    return [dict(zip(header, r)) for r in rows[1:] if r and any(v is not None for v in r)]


def validate():
    manifest = json.loads((GAP / 'manifest.json').read_text())
    expected = {i['item_id']: (i['cv_id'], i['job_id']) for i in manifest['items']}
    dev = set((ROOT / 'evals/splits/dev_job_ids.txt').read_text().split())
    base = [json.loads(l) for l in (BASE / 'relevance_gold.jsonl').read_text().splitlines() if l.strip()]
    existing = {(r['cv_id'], r['job_id']) for r in base}
    problems, labels = [], []
    rows = read_rows()
    seen = Counter(r.get('item_id') for r in rows)
    for item_id, n in seen.items():
        if item_id not in expected:
            problems.append(f'unexpected row {item_id}')
        elif n > 1:
            problems.append(f'duplicate row {item_id}')
    missing = sorted(set(expected) - set(seen))
    if missing:
        problems.append(f'missing rows {missing}')
    for r in rows:
        item = r.get('item_id')
        if item not in expected:
            continue
        cv, job = expected[item]
        if (r.get('cv_id'), r.get('job_id')) != (cv, job):
            problems.append(f'{item}: CV or job changed')
            continue
        value = r.get('relevance_0_3')
        try:
            value = int(value)
        except (TypeError, ValueError):
            problems.append(f'{item}: relevance empty or not a number')
            continue
        if value not in range(4):
            problems.append(f'{item}: relevance {value} outside 0..3')
        reason = (r.get('main_reason') or '').strip()
        if not reason:
            problems.append(f'{item}: main_reason empty')
        if job not in dev:
            problems.append(f'{item}: {job} is not a development job')
        if (cv, job) in existing:
            problems.append(f'{item}: {cv}/{job} already in r3')
        labels.append({'cv_id': cv, 'job_id': job, 'job_title': r.get('job_title'), 'relevance_0_3': value,
                       'main_reason': reason, 'constraint_note': (r.get('constraint_note') or None),
                       'label_source': 'annotator', 'review_status': 'approved', 'review_action': 'added',
                       'gold_review': 'yes', 'guideline_version': 'v1.3', 'split': 'development',
                       'labeling_mode': 'blind_no_rank_score_method_or_draft',
                       'approval_provenance': {'source': str(WORKBOOK.relative_to(ROOT)),
                                               'gap_manifest': str((GAP / 'manifest.json').relative_to(ROOT)),
                                               'decision': 'D-080', 'annotator': 'Dion'}})
    return labels, problems, base


def main(write: bool):
    labels, problems, base = validate()
    dist = Counter((l['cv_id'], l['relevance_0_3']) for l in labels)
    print(json.dumps({'rows_valid': len(labels), 'problems': problems,
                      'distribution': {f'{cv}:{v}': n for (cv, v), n in sorted(dist.items())},
                      'r3_relevance_rows': len(base), 'r4_relevance_rows_if_written': len(base) + len(labels)}, indent=1))
    if not write:
        return
    if problems:
        raise SystemExit('Fix the workbook first; nothing was written')
    if TARGET.exists():
        raise SystemExit('Bundle r4 already exists; make a new version instead')
    before = {p.name: digest(p) for p in BASE.iterdir() if p.is_file()}
    TARGET.mkdir(parents=True)
    for p in BASE.iterdir():
        if p.is_file() and p.name not in ('manifest.json', 'relevance_gold.jsonl'):
            shutil.copy2(p, TARGET / p.name)
    with (TARGET / 'relevance_gold.jsonl').open('x', encoding='utf-8') as f:
        f.write((BASE / 'relevance_gold.jsonl').read_text(encoding='utf-8'))
        for row in labels:
            f.write(json.dumps(row, ensure_ascii=False) + '\n')
    files = {p.name: digest(p) for p in TARGET.iterdir() if p.is_file()}
    manifest = {'schema_version': 'development-gold-amendment-v1', 'bundle': TARGET.name,
                'base_bundle': BASE.name, 'base_manifest_sha256': digest(BASE / 'manifest.json'),
                'decision': 'D-080', 'added_relevance_rows': len(labels),
                'relevance_rows': len(base) + len(labels), 'unchanged_files_copied_from_base': sorted(
                    n for n in files if n != 'relevance_gold.jsonl'),
                'workbook_sha256': digest(WORKBOOK), 'files': files, 'test_used': False}
    with (TARGET / 'manifest.json').open('x') as f:
        json.dump(manifest, f, indent=2)
    after = {p.name: digest(p) for p in BASE.iterdir() if p.is_file()}
    if before != after:
        raise SystemExit('Base bundle changed during import; investigate')
    print(json.dumps({'written': str(TARGET.relative_to(ROOT)), 'added': len(labels)}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--write', action='store_true')
    main(parser.parse_args().write)
