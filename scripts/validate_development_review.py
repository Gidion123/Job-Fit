"""Read-only draft QA: exact sources, split isolation, pending labels and review cap."""
import collections
import csv
import hashlib
import json
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from jobfit.config import REPO_ROOT
from openpyxl import load_workbook


def records(sheet):
    rows=list(sheet.values)
    return [dict(zip(rows[0],r))for r in rows[1:] if any(x is not None for x in r)]


def validate():
    root=REPO_ROOT
    if (root/'evals/labeling/JobFit_Development_Labeling_v0.1.xlsx').exists():
        from validate_combined_labeling import validate as validate_combined
        return validate_combined()
    payload=json.loads((root/'evals/labeling/drafts/development_review_v1.json').read_text())
    dev=set((root/'evals/splits/dev_job_ids.txt').read_text().splitlines())
    test=set((root/'evals/splits/test_job_ids.txt').read_text().splitlines())
    for name, digest in payload['source_hashes'].items():
        assert hashlib.sha256((root/name).read_bytes()).hexdigest()==digest, name
    manifest=json.loads((root/'evals/splits/split_manifest.json').read_text())
    for name,digest in manifest['output_hashes'].items():
        assert hashlib.sha256((root/'evals/splits'/name).read_bytes()).hexdigest()==digest
    pool=list(csv.DictReader((root/'evals/pools/dev_pool.csv').open()))
    expected={(r['cv_id'],r['job_id']):r for r in pool if r['already_gold']=='no'}
    books={i:load_workbook(root/f'evals/labeling/dev_batch_0{i}.xlsx')for i in [1,2]}
    audit={'status':'passed','api_calls':0,'session_cost_usd':0,'workbooks':{}}
    for i,w in books.items():
        sources=records(w['JDs'])
        assert {r['job_id'] for r in sources} <= dev
        assert not {r['job_id'] for r in sources}&test
        for r in sources:
            assert r['jd_text']==payload['jobs'][r['job_id']]['description_clean']
        cvs=records(w['CVs'])
        assert {r['cv_id']for r in cvs}=={'CV1','CV2'}
        assert all(r['cv_text']==payload['cvs'][r['cv_id']]for r in cvs)
        sheet=w['A_Extraction']if i==1 else w['C_Relevance']
        rows=records(sheet)
        assert all(r['review_status']=='pending' and r['review_action'] is None
                   and r['label_source']=='model_draft' and r['guideline_version']=='v1.2'for r in rows)
        assert len(sheet.tables)==1 and len(sheet.data_validations.dataValidation)>=5
        assert sheet.freeze_panes=='A2'
        draft_rows = payload['units'] if i == 1 else payload['relevance']
        lookup = {r['unit_no'] if i == 1 else (r['cv_id'], r['job_id']): r for r in draft_rows}
        for row in rows:
            key = row['unit_no'] if i == 1 else (row['cv_id'], row['job_id'])
            for field, value in row.items():
                assert (value if value is not None else '') == (lookup[key].get(field) if lookup[key].get(field) is not None else ''), (key, field)
        if i==1:
            assert len(rows)==84
            assert len({r['unit_no']for r in rows})==len(rows)
            assert all(r['source_quote'] in payload['jobs'][r['job_id']]['description_clean'] for r in rows)
            assert not records(w['B_Evidence'])
            counts=dict(collections.Counter(r['job_id']for r in rows))
        else:
            assert len(rows)==60 and len(expected)==60
            assert {(r['cv_id'],r['job_id'])for r in rows}==set(expected)
            assert all(isinstance(r['relevance_0_3'], int) and 0 <= r['relevance_0_3'] <= 3 for r in rows)
            assert all(r['gold_review']==expected[(r['cv_id'],r['job_id'])]['gold_review']for r in rows)
            assert all(r['gold_review']=='yes'for r in rows[:40])
            assert all(r['gold_review']=='no'for r in rows[40:])
            assert collections.Counter(r['cv_id']for r in rows[:40])=={'CV1':20,'CV2':20}
            assert all(expected[(r['cv_id'],r['job_id'])]['gold_review']=='yes'for r in pool
                       if r['already_gold']=='no'and r['in_top5_any']=='yes')
            assert not any(h in [c.value for c in sheet[1]]for h in ['score','best_rank','methods_top10'])
            counts={cv:dict(collections.Counter(r['relevance_0_3']for r in rows if r['cv_id']==cv))for cv in ['CV1','CV2']}
        cached=load_workbook(root/f'evals/labeling/dev_batch_0{i}.xlsx',data_only=True)
        assert all(cached['Timing'].cell(r,5).value in [None,'']for r in range(2,w['Timing'].max_row+1))
        for s in w:
            assert all(c.data_type!='e'for row in s for c in row)
        audit['workbooks'][f'dev_batch_0{i}.xlsx']={'sha256':hashlib.sha256((root/f'evals/labeling/dev_batch_0{i}.xlsx').read_bytes()).hexdigest(),
             'rows':len(rows),'counts':counts,'sheets':w.sheetnames,'sources_exact':True,'new_labels_pending':True}
    audit['protected_sources_unchanged']=True
    audit['all_top5_covered_within_40_new_reviews']=True
    return audit


if __name__=='__main__':
    audit=validate()
    name=('development_combined_workbook_qa_20261002.json' if 'workbook' in audit else 'development_workbook_qa_20261001.json')
    (REPO_ROOT/'evals/results'/name).write_text(json.dumps(audit,indent=2)+'\n')
    print(json.dumps(audit,indent=2))
