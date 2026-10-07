"""Import Dion's completed gap workbook (D-080, D-085) into gold bundle r4. Dry run unless --write.

Source: evals/labeling/dev_relevance_gap_20261004_v2/JobFit_Dev_Relevance_Gap_v2_v1.3_Final.xlsx
- Relevance: the Items sheet is Dion's answer sheet. Every value must be 0..3, except
  G04 (F00559), which Dion marked UNSCORABLE: it becomes a source-quality hold, NOT a 0
  (D-085, D-052: a missing judgment is never a negative).
- Extraction (A) and evidence (B) for the gap jobs go to separate files with their own
  unit namespace (`gap_v2:`). Seven of these JDs already have r3 extraction gold; the r3
  units stay authoritative and untouched, so the gap units never replace or mix with them.
- Provenance: the sheets carry a model draft accepted by Dion (draft_note "Stage 5 model
  draft"), so labeling_mode is recorded as model-draft-assisted, human-accepted.
Checks: workbook closed, identities match the manifest, relevance values, quotes found in
the JD/CV text, every B unit exists in A. r3 is copied byte for byte and re-hashed after.
"""
from __future__ import annotations

import argparse
from collections import Counter
from hashlib import sha256
import json
from pathlib import Path
import re
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'src')]

GAP = ROOT / 'evals/labeling/dev_relevance_gap_20261004_v2'
WORKBOOK = GAP / 'JobFit_Dev_Relevance_Gap_v2_v1.3_Final.xlsx'
BASE = ROOT / 'evals/gold/development_v13_reviewed_20261003_stage1_r3'
TARGET = ROOT / 'evals/gold/development_v13_reviewed_20261004_gap_r4'
UNIT_SET = 'gap_v2'
UNSCORABLE = {'UNSCORABLE'}
CVS = {'CV1': 'data/synthetic_cvs/cv_01_fresh_graduate_data_science_id.md',
       'CV2': 'data/synthetic_cvs/cv_02_career_switcher_ai_engineer_en.md'}
MODE = 'model_draft_assisted_human_accepted'


def digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def ws(s: str) -> str:
    return re.sub(r'\s+', ' ', s).strip()


def quote_ok(quote: str | None, text: str) -> bool:
    parts = [p for p in re.split(r'\s*\|\|\s*|\n\s*\.\.\.\s*\n', quote or '') if p.strip()]
    return bool(parts) and all(p in text or ws(p) in ws(text) for p in parts)


def sheets() -> dict[str, list[dict]]:
    from openpyxl import load_workbook
    if any(GAP.glob('~$*')):
        raise SystemExit('Close the workbook in Excel first (lock file present)')
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    out = {}
    for name in ('Items', 'A_Extraction', 'B_Evidence', 'C_Relevance'):
        rows = list(wb[name].iter_rows(values_only=True))
        out[name] = [dict(zip(rows[0], r)) for r in rows[1:] if r and any(v is not None for v in r)]
    return out


def build():
    from jobfit.cv.text_extract import extract_text
    s = sheets()
    manifest = json.loads((GAP / 'manifest.json').read_text())
    expected = {i['item_id']: (i['cv_id'], i['job_id']) for i in manifest['items']}
    dev = set((ROOT / 'evals/splits/dev_job_ids.txt').read_text().split())
    jd = {json.loads(l)['final_cluster_id']: json.loads(l)['description_clean']
          for l in (ROOT / 'data/processed/jobs_features.jsonl').read_text().splitlines()}
    cv_text = {k: extract_text((ROOT / v).read_bytes(), v).text for k, v in CVS.items()}
    base_rel = [json.loads(l) for l in (BASE / 'relevance_gold.jsonl').read_text().splitlines() if l.strip()]
    existing = {(r['cv_id'], r['job_id']) for r in base_rel}
    c_sheet = {(r['cv_id'], r['job_id']): r for r in s['C_Relevance']}
    provenance = {'source': str(WORKBOOK.relative_to(ROOT)), 'workbook_sha256': digest(WORKBOOK),
                  'gap_manifest': str((GAP / 'manifest.json').relative_to(ROOT)), 'decisions': ['D-080', 'D-085'],
                  'annotator': 'Dion', 'labeling_mode': MODE}
    problems, relevance, held = [], [], []
    seen = Counter(r['item_id'] for r in s['Items'])
    if set(seen) != set(expected) or any(n > 1 for n in seen.values()):
        problems.append('Items rows do not match the manifest exactly')
    for r in s['Items']:
        cv, job = expected.get(r['item_id'], (None, None))
        if (r['cv_id'], r['job_id']) != (cv, job):
            problems.append(f"{r['item_id']}: identity changed"); continue
        if job not in dev:
            problems.append(f'{r["item_id"]}: not a development job')
        if (cv, job) in existing:
            problems.append(f'{r["item_id"]}: already in r3')
        c = c_sheet.get((cv, job), {})
        raw = str(r['relevance_0_3']).strip() if r['relevance_0_3'] is not None else ''
        reason = (r.get('main_reason') or '').strip()
        if not reason:
            problems.append(f"{r['item_id']}: main_reason empty")
        if raw.upper() in UNSCORABLE:
            held.append({'record': 'C_Relevance', 'item_id': r['item_id'], 'cv_id': cv, 'job_id': job,
                         'hold_reason': 'unscorable_source_quality', 'annotator_value': raw,
                         'c_sheet_value': c.get('relevance_0_3'), 'main_reason': reason,
                         'source_note': r.get('source_note'), 'decision': 'D-085',
                         'note': 'Not a relevance judgment; never counted as 0 (D-052).'})
            continue
        if raw not in {'0', '1', '2', '3'}:
            problems.append(f"{r['item_id']}: relevance {raw!r} not in 0..3"); continue
        if c and str(c.get('relevance_0_3')) != raw:
            problems.append(f"{r['item_id']}: Items {raw} differs from C_Relevance {c.get('relevance_0_3')}")
        relevance.append({'cv_id': cv, 'cv_target': c.get('cv_target'), 'pilot_id': r['item_id'], 'job_id': job,
                          'job_title': r.get('job_title'), 'relevance_0_3': int(raw), 'main_reason': reason,
                          'constraint_note': r.get('constraint_note') or None, 'source_note': r.get('source_note') or None,
                          'c_sheet_reason': c.get('main_reason'), 'label_source': 'annotator',
                          'review_status': 'approved', 'review_action': c.get('review_action') or 'accepted',
                          'draft_note': c.get('draft_note'), 'review_note': None, 'guideline_version': 'v1.3',
                          'gold_review': 'yes', 'split': 'development', 'labeling_mode': MODE,
                          'approval_provenance': provenance})
    units = {}
    extraction = []
    for r in s['A_Extraction']:
        if r['job_id'] not in dev:
            problems.append(f"A {r['unit_no']}: not a development job")
        if not quote_ok(r['source_quote'], jd.get(r['job_id'], '')):
            problems.append(f"A {r['unit_no']}: source quote not in JD")
        if r['importance'] not in {'required', 'preferred', 'unknown'}:
            problems.append(f"A {r['unit_no']}: importance {r['importance']}")
        key = (r['job_id'], r['unit_no'])
        if key in units:
            problems.append(f"A {r['unit_no']}: duplicate")
        units[key] = r
        extraction.append({**r, 'unit_set': UNIT_SET, 'logical_unit_id': f"{UNIT_SET}:{r['job_id']}:{r['unit_no']}",
                           'split': 'development', 'labeling_mode': MODE, 'approval_provenance': provenance,
                           'jd_source_sha256': sha256(jd[r['job_id']].encode()).hexdigest()})
    evidence = []
    for r in s['B_Evidence']:
        key = (r['job_id'], r['unit_no'])
        if key not in units:
            problems.append(f"B {r['cv_id']} {r['unit_no']}: unit not in A"); continue
        if r['label'] not in {'MATCH', 'PARTIAL', 'NO_MATCH'}:
            problems.append(f"B {r['unit_no']}: label {r['label']}")
        if r['label'] in {'MATCH', 'PARTIAL'} and not quote_ok(r['cv_quote'], cv_text[r['cv_id']]):
            problems.append(f"B {r['cv_id']} {r['unit_no']}: CV quote not found")
        if r['check_status'] != 'done':
            held.append({'record': 'B_Evidence', 'cv_id': r['cv_id'], 'job_id': r['job_id'], 'unit_no': r['unit_no'],
                         'hold_reason': r['check_status'], 'label_in_sheet': r['label'], 'draft_note': r.get('draft_note'),
                         'decision': 'D-085'})
            continue
        evidence.append({**r, 'unit_set': UNIT_SET, 'logical_unit_id': f"{UNIT_SET}:{r['job_id']}:{r['unit_no']}",
                         'split': 'development', 'labeling_mode': MODE, 'approval_provenance': provenance})
    overlap = sorted({r['job_id'] for r in extraction} & {json.loads(l)['job_id'] for l in
                     (BASE / 'extraction_gold.jsonl').read_text().splitlines()})
    return {'relevance': relevance, 'extraction': extraction, 'evidence': evidence, 'held': held,
            'problems': problems, 'base_rel': base_rel, 'overlap_jobs_with_r3_extraction': overlap}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--write', action='store_true')
    args = ap.parse_args(argv)
    b = build()
    dist = Counter((r['cv_id'], r['relevance_0_3']) for r in b['relevance'])
    print(json.dumps({'relevance_rows': len(b['relevance']), 'held': [(h['record'], h.get('item_id') or h.get('unit_no'),
                      h['hold_reason']) for h in b['held']], 'extraction_units': len(b['extraction']),
                      'extraction_jobs': len({r['job_id'] for r in b['extraction']}),
                      'evidence_rows': len(b['evidence']),
                      'evidence_pairs': len({(r['cv_id'], r['job_id']) for r in b['evidence']}),
                      'overlap_jobs_with_r3_extraction (kept separate)': b['overlap_jobs_with_r3_extraction'],
                      'distribution': {f'{cv}:{v}': n for (cv, v), n in sorted(dist.items())},
                      'problems': b['problems']}, indent=1, ensure_ascii=False))
    if not args.write:
        return 0
    if b['problems']:
        raise SystemExit('Problems found; nothing written')
    if TARGET.exists():
        raise SystemExit('r4 already exists; make a new version')
    before = {p.name: digest(p) for p in BASE.iterdir() if p.is_file()}
    TARGET.mkdir(parents=True)
    for p in BASE.iterdir():
        if p.is_file() and p.name not in ('manifest.json', 'relevance_gold.jsonl'):
            shutil.copy2(p, TARGET / p.name)
    with (TARGET / 'relevance_gold.jsonl').open('x', encoding='utf-8') as f:
        f.write((BASE / 'relevance_gold.jsonl').read_text(encoding='utf-8'))
        for row in b['relevance']:
            f.write(json.dumps(row, ensure_ascii=False) + '\n')
    for name, rows in (('extraction_gold_gap_v2.jsonl', b['extraction']), ('evidence_gold_gap_v2.jsonl', b['evidence'])):
        with (TARGET / name).open('x', encoding='utf-8') as f:
            for row in rows:
                f.write(json.dumps(row, ensure_ascii=False, default=str) + '\n')
    (TARGET / 'held_gap_v2.json').write_text(json.dumps(b['held'], indent=1, ensure_ascii=False) + '\n')
    files = {p.name: digest(p) for p in TARGET.iterdir() if p.is_file()}
    manifest = {'schema_version': 'development-gold-amendment-v2', 'bundle': TARGET.name, 'base_bundle': BASE.name,
                'base_manifest_sha256': digest(BASE / 'manifest.json'), 'decisions': ['D-080', 'D-085'],
                'added_relevance_rows': len(b['relevance']), 'relevance_rows': len(b['base_rel']) + len(b['relevance']),
                'gap_unit_set': UNIT_SET, 'gap_extraction_units': len(b['extraction']),
                'gap_evidence_rows': len(b['evidence']), 'held_records': len(b['held']),
                'overlap_jobs_with_r3_extraction_kept_separate': b['overlap_jobs_with_r3_extraction'],
                'labeling_mode': MODE, 'workbook_sha256': digest(WORKBOOK), 'files': files, 'test_used': False,
                'unchanged_files_copied_from_base': sorted(n for n in before if n not in ('manifest.json', 'relevance_gold.jsonl'))}
    with (TARGET / 'manifest.json').open('x') as f:
        json.dump(manifest, f, indent=2)
    if before != {p.name: digest(p) for p in BASE.iterdir() if p.is_file()}:
        raise SystemExit('Base bundle changed during import; investigate')
    print('written', TARGET.relative_to(ROOT))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
