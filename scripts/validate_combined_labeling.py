"""Read-only QA of the unified development review workbook. Never exports gold."""
from collections import Counter
import hashlib
import json
from pathlib import Path
import sys
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
BOOK=ROOT/'evals/labeling/JobFit_Development_Labeling_v0.1.xlsx'
MANIFEST=ROOT/'evals/labeling/drafts/development_combined_manifest_v1.json'

def records(sheet):
    rows=list(sheet.values)
    return [dict(zip(rows[0],r)) for r in rows[1:] if any(v is not None for v in r)]

def fingerprint(row):
    return hashlib.sha256(json.dumps({k:row.get(k) or None for k in ['job_id','unit_text','source_quote','importance','category','group_id','min_years']},sort_keys=True).encode()).hexdigest()

def validate():
    m=json.loads(MANIFEST.read_text()); w=load_workbook(BOOK)
    assert w.sheetnames==list(m['headers']),w.sheetnames
    for name,h in m['headers'].items():assert list(next(w[name].values))==h,(name,'headers')
    sources={r['job_id']:r for r in records(w['JDs'])}
    cvs={r['cv_id']:r['cv_text'] for r in records(w['CVs'])}
    dev=set((ROOT/'evals/splits/dev_job_ids.txt').read_text().splitlines())
    test=set((ROOT/'evals/splits/test_job_ids.txt').read_text().splitlines())
    assert set(sources)==set(m['scope']['job_ids'])
    assert set(sources)<=dev and not set(sources)&test
    assert set(cvs)=={'CV1','CV2'}
    corpus={r['final_cluster_id']:r for r in map(json.loads,(ROOT/'data/processed/jobs_features.jsonl').read_text().splitlines())}
    assert all(s['jd_text']==corpus[j]['description_clean'] for j,s in sources.items())
    for cv,text in cvs.items():assert hashlib.sha256(text.encode()).hexdigest()==m['cv_body_sha256'][cv]
    a=records(w['A_Extraction']); b=records(w['B_Evidence']); c=records(w['C_Relevance'])
    units={r['unit_no']:r for r in a}; assert len(units)==len(a)
    assert all(r['source_quote'] in sources[r['job_id']]['jd_text'] for r in a)
    expected=set(map(tuple,m['scope']['pairs']))
    assert len(c)==len(expected) and {(r['cv_id'],r['job_id']) for r in c}==expected
    assert all(type(r['relevance_0_3']) is int and 0<=r['relevance_0_3']<=3 and r['main_reason'] for r in c)
    seen=set()
    for r in b:
        key=(r['cv_id'],r['job_id'],r['unit_no']); assert key not in seen;seen.add(key)
        u=units[r['unit_no']]; assert u['job_id']==r['job_id'] and u['unit_text']==r['unit_text'],key
        assert (r['cv_id'],r['job_id']) in expected
        assert r['label'] in ('MATCH','PARTIAL','NO_MATCH') and r['check_status'] in ('done','needs_clarification','failed')
        if r['label']=='NO_MATCH':assert not r['cv_quote'],key
        else:assert r['cv_quote'] and r['cv_quote'] in cvs[r['cv_id']],key
    zero=set(m['scope']['zero_requirement_jobs'])
    for cv,j in expected:
        ids={r['unit_no'] for r in a if r['job_id']==j}
        actual={r['unit_no'] for r in b if (r['cv_id'],r['job_id'])==(cv,j)}
        assert actual==ids,(cv,j,'incomplete evidence')
        assert bool(ids) or j in zero,(j,'unexpected zero requirements')
    preserved=m['preserved_rows']
    maps={'A_Extraction':units,'B_Evidence':{f'{r["cv_id"]}/{r["job_id"]}/{r["unit_no"]}':r for r in b},
          'C_Relevance':{f'{r["cv_id"]}/{r["job_id"]}':r for r in c}}
    for name,saved in preserved.items():
        for key,old in saved.items():
            current=maps[name][key]
            for field,value in old.items():assert (current.get(field) or None)==(value or None),(name,key,field,'changed preserved row')
    approved=Counter()
    for name,rows in [('A_Extraction',a),('B_Evidence',b),('C_Relevance',c)]:
        protected=set(preserved[name])
        for r in rows:
            key=r['unit_no'] if name=='A_Extraction' else (f'{r["cv_id"]}/{r["job_id"]}/{r["unit_no"]}' if name=='B_Evidence' else f'{r["cv_id"]}/{r["job_id"]}')
            assert r['review_status'] in ('approved','pending')
            if r['review_status']=='approved':
                approved[name]+=1
                assert r['review_action'] in ('accepted','edited','rejected','added')
                if r['review_action'] in ('edited','rejected'):assert r['review_note']
            if key not in protected:
                assert r['review_status']=='pending' and r['review_action'] is None and r['review_note'] is None
                assert r['label_source']=='model_draft' and r['guideline_version']=='v1.2'
    assert Counter(r['cv_id'] for r in c if r['gold_review']=='yes')=={'CV1':20,'CV2':20}
    assert {(r['cv_id'],r['job_id']) for r in c if r['gold_review']=='yes'}==set(map(tuple,m['mandatory_relevance_pairs']))
    stale=[r['unit_no'] for r in a if fingerprint(r)!=m['unit_fingerprints'][r['unit_no']]]
    assert not stale,('A changed; dependent B/C must be reviewed',stale)
    for name in ('A_Extraction','B_Evidence','C_Relevance'):
        s=w[name]; assert s.freeze_panes=='A2' and len(s.tables)==1
        validations={str(d.sqref):d for d in s.data_validations.dataValidation}
        header=list(next(s.values));col=header.index('review_action')+1
        from openpyxl.utils import get_column_letter
        dv=next(d for target,d in validations.items() if target.startswith(f'{get_column_letter(col)}2:'))
        assert dv.type=='list' and dv.formula1=='"accepted,edited,rejected,added"'
        assert dv.allow_blank and dv.showDropDown is False
    for s in w:
        assert all(cell.data_type!='e' for row in s for cell in row),(s.title,'formula error')
    cached=load_workbook(BOOK,data_only=True)
    for row in range(2,w['Timing'].max_row+1):
        if not w['Timing'].cell(row,3).value or not w['Timing'].cell(row,4).value:
            assert cached['Timing'].cell(row,5).value in (None,'')
    for rel,expected_hash in m['source_hashes'].items():
        p=ROOT/rel
        if not p.exists() and rel in m.get('archived_inputs',{}):p=ROOT/m['archived_inputs'][rel]
        assert p.exists() and hashlib.sha256(p.read_bytes()).hexdigest()==expected_hash,rel
    cases=[
      ('CV1','F00103','Data-quality frameworks','PARTIAL','done'),
      ('CV2','F00103','Data-quality frameworks','PARTIAL','done'),
      ('CV2','F00074','Direct API-level experience with an LLM provider (Anthropic | OpenAI | Google), not wrappers only','PARTIAL','done'),
      ('CV2','F00034','AI/ML framework: TensorFlow | PyTorch | scikit-learn | similar','PARTIAL','done'),
      ('CV2','F00114','LangChain | LangGraph | LlamaIndex | LangSmith | CrewAI | n8n agent framework','PARTIAL','done'),
      ('CV2','F00798','Kubernetes | OpenShift | containers | cloud/enterprise production environment','PARTIAL','done'),
      ('CV2','F00208','AI/ML/cloud certification','NO_MATCH','done'),
      ('CV1','F00556','Currently pursuing a quantitative/technical degree','PARTIAL','done'),
      ('CV1','F00369','NLP sequence models (RNN/LSTM with TensorFlow | PyTorch)','PARTIAL','done'),
      ('CV2','F00029','PyTorch | Hugging Face | scikit-learn','MATCH','done'),
      ('CV1','F00208','Technical/quantitative degree | equivalent education and experience','MATCH','done'),
      ('CV1','F00090','At least 2-4 years backend/software work focused on Python','PARTIAL','done'),
      ('CV2','F00003','At least 1 year in Generative AI projects','PARTIAL','done'),
    ]
    for cv,j,text,label,status in cases:
        row=next(r for r in b if (r['cv_id'],r['job_id'],r['unit_text'])==(cv,j,text))
        assert (row['label'],row['check_status'])==(label,status),(cv,j,text,'semantic spot check')
    return dict(status='passed',workbook=str(BOOK.relative_to(ROOT)),workbook_sha256=hashlib.sha256(BOOK.read_bytes()).hexdigest(),counts=dict(jds=len(sources),pairs=len(expected),a=len(a),b=len(b),c=len(c)),
      approved=dict(approved),pending=dict(A_Extraction=len(a)-approved['A_Extraction'],B_Evidence=len(b)-approved['B_Evidence'],C_Relevance=len(c)-approved['C_Relevance']),
      zero_requirement_jobs=sorted(zero),quotes_exact=True,sources_unchanged=True,preserved_rows_unchanged=True,
      full_pair_unit_coverage=True,split_isolation=True,review_validation=True,mandatory_review_pairs=40,
      semantic_spot_checks=len(cases),api_calls=0,session_cost_usd=0,limits=['Draft semantic correctness still requires annotator review.','Native Excel interaction not exercised; saved XML validation and artifact previews verified.'])

if __name__=='__main__':
    result=validate()
    (ROOT/'evals/results/development_combined_workbook_qa_20261002.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result,indent=2))
