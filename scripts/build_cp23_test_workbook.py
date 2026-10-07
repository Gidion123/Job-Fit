"""Build the D-053 blind test relevance pool and workbook from a saved test run.

Input: a test run JSON with, per CV, `stage1_ids` (original stage-1 order) and
`final_order` (product order of the analyzed top K). The run must point to an
APPROVED freeze receipt with no file drift.

Default is a dry run: it prints pair counts, overlap and the time estimate that
D-053 asks for before any workbook exists. `--write` creates:
- evals/pools/test_pool_<run>_v<n>/pool_provenance.json  (ranks, hashes, run id)
- evals/labeling/test_relevance_<run>_v<n>/JobFit_Test_Relevance.xlsx (blind)
- evals/labeling/test_relevance_<run>_v<n>/manifest.json (items only, no ranks)
No model call, no DB, no label is read.
"""
from __future__ import annotations

import argparse
from datetime import date
from hashlib import sha256
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'src')]

from jobfit.eval.test_pool import blind_items, build_pool
from scripts.prepare_cp23_freeze import CVS, verify

SEED = 20261053
FEATURES = ROOT / 'data/processed/jobs_features.jsonl'
TEST_IDS = ROOT / 'evals/splits/test_job_ids.txt'
README = [
    'JobFit test relevance labels (D-053). Blind held-out pool.',
    'Scope: union of the original top 10 stage-1 jobs and the top 10 final recommendations, per CV, test jobs only.',
    'Blind: no rank, score, method, model suggestion or experience bucket is shown. Rows are shuffled.',
    'Fill relevance_0_3 (0,1,2,3) and main_reason. constraint_note is optional.',
    'Rule: annotation guideline v1.3 Part D. Judge in automatic mode (four target role families).',
    '3 = most required technical requirements supported by CV evidence and no explicit constraint conflict',
    '2 = several key required requirements supported, some gaps, no explicit conflict',
    '1 = weak evidence support, OR an explicit conflict (for example 3+ years asked, CV shows under 1 year)',
    '0 = outside the four target role families, or almost no requirement supported',
    'Soft skills and location do not lower relevance by themselves.',
    'Labels made here are never used to change the model, prompt, K, weight or rules (D-046, D-053).',
    'Save, close Excel and tell Claude. Nothing becomes gold until you approve the import.',
]
COLUMNS = ['item_id', 'cv_id', 'job_id', 'job_title', 'company', 'job_description',
           'relevance_0_3', 'main_reason', 'constraint_note', 'source_note']


def load_run(path: Path) -> dict:
    run = json.loads(path.read_text())
    freeze = ROOT / run['freeze_folder']
    approved = freeze / 'freeze_receipt_APPROVED.json'
    if not approved.exists():
        raise SystemExit('The test run does not point to an approved freeze receipt')
    if json.loads(approved.read_text())['status'] != 'APPROVED':
        raise SystemExit('Freeze receipt is not approved')
    drift = verify(freeze)
    if drift:
        raise SystemExit(f'Files changed since the freeze: {drift}')
    return run


def next_version(stem: str) -> int:
    n = 1
    while (ROOT / f'evals/labeling/{stem}_v{n}').exists() or (ROOT / f'evals/pools/{stem.replace("test_relevance", "test_pool")}_v{n}').exists():
        n += 1
    return n


def write_workbook(path: Path, items: list[dict], jobs: dict, cv_ids: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = 'README'
    for line in README:
        ws.append([line])
    ws.column_dimensions['A'].width = 140
    cvs = wb.create_sheet('CVs')
    cvs.append(['cv_id', 'cv_text'])
    for cv in cv_ids:
        cvs.append([cv, (ROOT / CVS[cv]).read_text()])
    sheet = wb.create_sheet('Items')
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for it in items:
        job = jobs[it['job_id']]
        sheet.append([it['item_id'], it['cv_id'], it['job_id'], job['title'], job['company'],
                      job['description_clean'], None, None, None, None])
    dv = DataValidation(type='whole', operator='between', formula1='0', formula2='3', allow_blank=True)
    sheet.add_data_validation(dv)
    dv.add(f'G2:G{len(items) + 1}')
    for col, width in zip('ABCDEFGHIJ', (8, 6, 9, 40, 28, 90, 12, 40, 30, 30)):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.freeze_panes = 'D2'
    wb.save(path)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('run', type=Path, help='saved test run JSON')
    ap.add_argument('--write', action='store_true')
    args = ap.parse_args(argv)
    run = load_run(args.run)
    eligible = set(TEST_IDS.read_text().split())
    stage1 = {cv: r['stage1_ids'] for cv, r in run['cvs'].items()}
    final = {cv: r['final_order'] for cv, r in run['cvs'].items()}
    if not set(stage1) <= set(CVS):
        raise SystemExit('Unknown CV id in the run')
    pool = build_pool(stage1, final, eligible=eligible)
    print(json.dumps({'run_id': run['run_id'], 'pairs': pool['pairs'], 'per_cv': pool['per_cv'],
                      'effort': pool['effort']}, indent=1))
    if not args.write:
        print('Dry run. Share these counts with Dion and agree on batching before --write (D-053).')
        return 0
    jobs = {}
    for line in FEATURES.read_text().splitlines():
        row = json.loads(line)
        jobs[row['final_cluster_id']] = row
    items = blind_items(pool['rows'], seed=SEED)
    stem = f"test_relevance_{run['run_id']}"
    n = next_version(stem)
    lab = ROOT / f'evals/labeling/{stem}_v{n}'
    prov = ROOT / f"evals/pools/test_pool_{run['run_id']}_v{n}"
    lab.mkdir(parents=True)
    prov.mkdir(parents=True)
    book = lab / 'JobFit_Test_Relevance.xlsx'
    write_workbook(book, items, jobs, sorted(stage1))
    (lab / 'manifest.json').write_text(json.dumps({
        'version': f'test-relevance-{run["run_id"]}-v{n}', 'date': date.today().isoformat(), 'split': 'test',
        'rule': pool['rule'], 'items': items, 'blind': True,
        'hidden': ['rank', 'score', 'method', 'model suggestion', 'experience bucket'],
        'shuffle_seed': SEED, 'workbook_sha256_at_creation': sha256(book.read_bytes()).hexdigest()}, indent=1) + '\n')
    (prov / 'pool_provenance.json').write_text(json.dumps({
        'run_id': run['run_id'], 'run_file': str(args.run), 'run_sha256': sha256(args.run.read_bytes()).hexdigest(),
        'freeze_folder': run['freeze_folder'], 'test_ids_sha256': sha256(TEST_IDS.read_bytes()).hexdigest(),
        **pool}, indent=1) + '\n')
    print('wrote', lab.relative_to(ROOT), 'and', prov.relative_to(ROOT))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
