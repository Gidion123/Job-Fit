"""Export the approved pilot labels to evals/gold/ as the development split (D-038 step 5).

Reads evals/pilot/JobFit_Pilot_Labeling_v0.1.xlsx and writes only rows with
review_status = approved and review_action != rejected. Rows that are still pending
stop the export, so a half-reviewed batch never becomes gold.

Usage: python scripts/export_pilot_gold.py
"""
from __future__ import annotations

import csv
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "evals" / "pilot" / "JobFit_Pilot_Labeling_v0.1.xlsx"
GOLD = ROOT / "evals" / "gold"
SPLITS = ROOT / "evals" / "splits"
SPLIT = "development"
SNAPSHOT_ID = "CP1_20260926"


def _rows(ws, key: str) -> list[dict]:
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    return [r for r in rows if r.get(key) not in (None, "")]


def _clean(value):
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _num(value):
    return None if value in (None, "") else float(value)


def _quote_ok(quote: str | None, text: str) -> bool:
    if not quote:
        return True
    norm = lambda s: " ".join(s.split())
    return all(norm(part.rstrip(".")) in norm(text) for part in quote.split(" | ") if part.strip())


def main() -> int:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    jds = {r["job_id"]: r for r in _rows(wb["JDs"], "job_id")}
    cvs = {r["cv_id"]: r for r in _rows(wb["CVs"], "cv_id")}
    a = _rows(wb["A_Extraction"], "unit_text")
    b = _rows(wb["B_Evidence"], "unit_text")
    c = _rows(wb["C_Relevance"], "relevance_0_3")

    pending = [r for r in a + b + c if r.get("review_status") != "approved"]
    if pending:
        print(f"Stopped: {len(pending)} rows are not approved yet.", file=sys.stderr)
        return 1

    common = {"split": SPLIT, "snapshot_id": SNAPSHOT_ID, "review_status": "approved", "reviewed_by": "annotator"}
    bad_quotes = []

    GOLD.mkdir(parents=True, exist_ok=True)
    with open(GOLD / "extraction_gold.jsonl", "w", encoding="utf-8") as f:
        for r in a:
            if r["review_action"] == "rejected":
                continue
            if not _quote_ok(r["source_quote"], jds[r["job_id"]]["jd_text"]):
                bad_quotes.append(r["unit_no"])
            rec = {**common, "pilot_id": r["pilot_id"], "job_id": r["job_id"], "unit_no": r["unit_no"],
                   "unit_text": _clean(r["unit_text"]), "source_quote": _clean(r["source_quote"]),
                   "importance": r["importance"], "category": r["category"], "group_id": _clean(r["group_id"]),
                   "min_years": _num(r["min_years"]), "label_source": r["label_source"],
                   "review_action": r["review_action"], "review_note": _clean(r["review_note"]),
                   "guideline_version": r["guideline_version"]}
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    with open(GOLD / "evidence_gold.jsonl", "w", encoding="utf-8") as f:
        for r in b:
            if r["review_action"] == "rejected":
                continue
            if not _quote_ok(r["cv_quote"], cvs[r["cv_id"]]["cv_text"]):
                bad_quotes.append(f'{r["cv_id"]} x {r["unit_no"]}')
            rec = {**common, "cv_id": r["cv_id"], "job_id": r["job_id"], "unit_no": r["unit_no"],
                   "label": r["label"], "check_status": r["check_status"], "cv_quote": _clean(r["cv_quote"]),
                   "cv_section": _clean(r["cv_section"]), "label_source": r["label_source"],
                   "review_action": r["review_action"], "review_note": _clean(r["review_note"]),
                   "guideline_version": r["guideline_version"]}
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    fields = ["split", "snapshot_id", "cv_id", "pilot_id", "job_id", "relevance", "label_source",
              "review_action", "guideline_version"]
    with open(GOLD / "relevance_gold.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in c:
            if r["review_action"] == "rejected":
                continue
            w.writerow({"split": SPLIT, "snapshot_id": SNAPSHOT_ID, "cv_id": r["cv_id"], "pilot_id": r["pilot_id"],
                        "job_id": r["job_id"], "relevance": int(r["relevance_0_3"]), "label_source": r["label_source"],
                        "review_action": r["review_action"], "guideline_version": r["guideline_version"]})

    SPLITS.mkdir(parents=True, exist_ok=True)
    dev_ids = sorted({r["job_id"] for r in a} | {r["job_id"] for r in c})
    (SPLITS / "dev_job_ids.txt").write_text("\n".join(dev_ids) + "\n", encoding="utf-8")

    digest = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()[:16]
    model_rows = [r for r in a + b + c if r["label_source"] == "model_draft"]
    summary = {
        "source_workbook_sha256_16": digest,
        "extraction_units": len(a), "evidence_rows": len(b), "relevance_labels": len(c),
        "dev_job_ids": dev_ids,
        "review_actions": dict(Counter(r["review_action"] for r in a + b + c)),
        "acceptance_rate_model_drafts": round(sum(r["review_action"] == "accepted" for r in model_rows) / len(model_rows), 3),
        "quotes_not_found": bad_quotes,
    }
    print(json.dumps(summary, indent=2))
    return 0 if not bad_quotes else 2


if __name__ == "__main__":
    raise SystemExit(main())
